"""Renderer capability fixtures. Run only inside the document sandbox, not on the host."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4

root = Path('/vercel/sandbox')
root.mkdir(exist_ok=True)
out = root / 'fixtures'
out.mkdir(exist_ok=True)
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
section.header_distance = Inches(.5)
normal = doc.styles['Normal']
normal.font.name, normal.font.size = 'Times New Roman', Pt(12)
normal.paragraph_format.line_spacing = 2
normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.add_run('Rivera ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); header._p.append(field)
for line in ['Alex Rivera', 'Professor Chen', 'English 101', '24 October 2026']:
    doc.add_paragraph(line)
p = doc.add_paragraph('Learning Through Ensemble Performance'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for text in [
    'An ensemble performance depends on a balance between individual expression and collective responsibility. A musician can introduce an idea, but its effect depends on how the other performers respond. Listening therefore becomes an active part of the composition rather than a pause between individual contributions.',
    'The relationship between repetition and variation makes this process audible. When an opening phrase returns with a different rhythm or register, the audience recognizes a familiar idea while hearing its meaning change. The difference is created through decisions made in response to the surrounding performance.',
    'This pattern also offers a useful model for collaborative learning. Participants need enough freedom to explore, but they must remain attentive to a shared purpose. The quality of the result depends on both the strength of individual contributions and the care with which those contributions are connected.'
]:
    p = doc.add_paragraph(text); p.paragraph_format.first_line_indent = Inches(.5)
p = doc.add_paragraph('Works Cited'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.page_break_before = True
p = doc.add_paragraph('Sample Author. A Demonstration Source for Formatting Tests. Example Press, 2026. This is a fictional fixture, not a real citation.')
p.paragraph_format.left_indent = Inches(.5); p.paragraph_format.first_line_indent = Inches(-.5)
doc.save(out / 'mla.docx')
assert all(p.paragraph_format.first_line_indent == Inches(.5) for p in doc.paragraphs[5:8])

# Branded report, image, table, header/footer and a targeted revision.
im = Image.new('RGB', (800, 240), '#e9edf6'); draw = ImageDraw.Draw(im)
draw.rectangle((30, 35, 770, 205), fill='#26456c'); im.save(out / 'banner.png')
report = Document(); report.add_heading('Quarterly Project Review', 0)
report.sections[0].header.paragraphs[0].text = 'RESEARCH STUDIO | INTERNAL'
report.sections[0].footer.paragraphs[0].text = 'Prepared for the project team'
report.add_picture(str(out / 'banner.png'), width=Inches(6))
report.add_paragraph('The team completed three milestones. The next phase focuses on validation and a controlled launch.')
table = report.add_table(rows=1, cols=3); table.style = 'Light Shading Accent 1'
for cell, text in zip(table.rows[0].cells, ['Milestone', 'Owner', 'Status']): cell.text = text
for row in [('Research', 'Alex', 'Complete'), ('Prototype', 'Sam', 'Complete'), ('Validation', 'Jo', 'In progress')]:
    for cell, text in zip(table.add_row().cells, row): cell.text = text
report.add_heading('Next steps', 1); report.add_paragraph('Validate the prototype with five reviewers and document the findings.')
report.save(out / 'report.docx')
revised = Document(out / 'report.docx'); revised.paragraphs[0].text = 'Updated Project Review'; revised.save(out / 'report-revised.docx')
assert len(revised.inline_shapes) == 1 and len(revised.tables) == 1
assert revised.sections[0].header.paragraphs[0].text == report.sections[0].header.paragraphs[0].text

# Direct PDF path with deliberate page breaks and vector content.
c = canvas.Canvas(str(out / 'brief.pdf'), pagesize=A4)
for n in range(1, 3):
    c.setFillColorRGB(.15, .27, .42); c.rect(0, 740, 600, 105, fill=1, stroke=0)
    c.setFillColorRGB(1, 1, 1); c.setFont('Helvetica-Bold', 23); c.drawString(45, 780, 'Project Brief')
    c.setFillColorRGB(.1, .1, .1); c.setFont('Helvetica', 12)
    c.drawString(45, 690, 'A clear objective, measurable milestones, and a practical next step.')
    c.drawString(45, 665, 'The layout is authored directly with a mature PDF library.')
    c.setFont('Helvetica', 10); c.drawRightString(550, 35, f'Page {n}'); c.showPage()
c.save()
print('Fixtures written: MLA essay, report, preserved revision, two-page PDF')
